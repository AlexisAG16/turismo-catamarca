from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "informe-general-turismo-catamarca-2026.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(35, 35, 35)
MUTED = RGBColor(89, 89, 89)
HEADER_FILL = "E8EEF5"
CALLOUT_FILL = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)

    for tag, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_cell_margins(table)
    for row in table.rows:
        for idx, width in enumerate(widths):
            cell = row.cells[idx]
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Página ")
    run.font.size = Pt(9)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.add_run(text)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Body Text"
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    return p


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.style = "Body Text"
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    p.add_run(f" {text}")
    doc.add_paragraph()


def add_label_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.8, 4.7])
    hdr = table.rows[0].cells
    hdr[0].text = "Aspecto"
    hdr[1].text = "Detalle vigente"
    for cell in hdr:
        set_cell_shading(cell, HEADER_FILL)
        for p in cell.paragraphs:
            p.runs[0].bold = True
            p.runs[0].font.color.rgb = DARK_BLUE
    for label, detail in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = detail
        cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cells[1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    doc.add_paragraph()
    return table


def configure_styles(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    body = doc.styles["Body Text"]
    body.font.name = "Calibri"
    body.font.size = Pt(11)
    body.font.color.rgb = INK
    body.paragraph_format.space_after = Pt(6)
    body.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.25


def build():
    doc = Document()
    configure_styles(doc)

    header = doc.sections[0].header.paragraphs[0]
    header.text = "Turismo Catamarca - Informe de estudio"
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = MUTED

    footer = doc.sections[0].footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Informe de Especificaciones y Contexto Teórico")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = DARK_BLUE

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Proyecto Turismo Catamarca")
    r.font.size = Pt(14)
    r.font.color.rgb = MUTED

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(f"Documento de estudio generado el {date.today().strftime('%d/%m/%Y')}")
    r.font.size = Pt(10)
    r.font.color.rgb = MUTED

    add_callout(
        doc,
        "Propósito.",
        "Este documento resume el problema, la solución, las funcionalidades vigentes y las decisiones técnicas del proyecto para facilitar el estudio y la defensa conceptual del sistema.",
    )

    add_heading(doc, "1. Contexto Del Problema", 1)
    add_body(
        doc,
        "La información turística de Catamarca suele encontrarse dispersa entre sitios institucionales, redes sociales, mapas, videos, notas periodísticas y recomendaciones informales. Esto genera una experiencia fragmentada para turistas y residentes: descubrir un lugar, verificar su ubicación, conocer actividades y armar un recorrido exige consultar varias fuentes.",
    )
    add_body(
        doc,
        "El proyecto propone una plataforma web centralizada para organizar atractivos, circuitos y actividades de la provincia. La solución aporta estructura, navegación, filtros, multimedia, mapas e itinerario personal, transformando datos dispersos en una experiencia turística consultable y administrable.",
    )

    add_heading(doc, "2. Destinatarios Y Contexto De Uso", 1)
    for item in [
        "Turistas que planifican una visita a Catamarca y necesitan información confiable.",
        "Residentes que desean redescubrir atractivos por departamento.",
        "Usuarios registrados que guardan lugares de interés en un itinerario personal.",
        "Administradores que cargan, corrigen y mantienen datos turísticos.",
        "Evaluadores o docentes que necesitan revisar alcance, dominio y arquitectura.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3. Objetivo De La Solución", 1)
    add_body(
        doc,
        "El objetivo principal es centralizar la oferta turística de Catamarca en una aplicación web responsive, conectada a MongoDB y con herramientas administrativas. El sistema permite explorar atractivos, consultar circuitos, revisar actividades, abrir ubicaciones, acceder a videos y guardar un itinerario.",
    )

    add_heading(doc, "4. Alcance Funcional MoSCoW", 1)
    add_label_table(
        doc,
        [
            ("Must Have", "Atractivos, circuitos, actividades, autenticación, rol administrador, formularios, validación backend, paginación, mapas, videos, imágenes con fallback e itinerario."),
            ("Should Have", "Detalle de circuito, estados vacíos, informes Excel, modo claro/oscuro, botón volver conservando paginado, confirmaciones visuales y SEO básico."),
            ("Could Have", "Categorías, dificultad de actividades, carga de imágenes propia, favoritos sincronizados, recomendaciones y métricas."),
            ("Won't Have", "Reservas, pagos, contratación directa de servicios, chat interno, geolocalización en tiempo real y app móvil nativa."),
        ],
    )

    add_heading(doc, "5. Funcionalidades Vigentes", 1)
    add_heading(doc, "5.1 Atractivos", 2)
    add_body(
        doc,
        "El módulo de atractivos es el núcleo del sitio. Presenta tarjetas con imagen, departamento, nombre, descripción breve y botones de mapa, video e itinerario. Soporta filtros por nombre, departamento y circuito, además de paginación backend de seis registros por página.",
    )
    add_heading(doc, "5.2 Detalle De Atractivo", 2)
    add_body(
        doc,
        "Cada atractivo tiene una página de detalle con imagen principal, descripción completa, departamento, actividades asociadas, mapa, video y botón de regreso que preserva la página de origen del paginado.",
    )
    add_heading(doc, "5.3 Circuitos", 2)
    add_body(
        doc,
        "Los circuitos agrupan atractivos en propuestas de recorrido. Se muestran como lista vertical de ancho completo y cuentan con página de detalle, donde se visualizan atractivos asociados y actividades derivadas.",
    )
    add_heading(doc, "5.4 Actividades", 2)
    add_body(
        doc,
        "Las actividades representan experiencias concretas vinculadas a atractivos. La relación vigente prioriza que el atractivo tenga una o varias actividades, en lugar de depender directamente del circuito.",
    )
    add_heading(doc, "5.5 Itinerario", 2)
    add_body(
        doc,
        "El itinerario permite guardar atractivos de interés, visualizar un contador, borrar elementos individuales, vaciar la lista completa y consultar una página dedicada con resumen y opción de impresión.",
    )
    add_heading(doc, "5.6 Administración E Informes", 2)
    add_body(
        doc,
        "El rol administrador habilita formularios de carga, edición y eliminación. También permite descargar un informe Excel desde soporte, con hojas para circuitos, atractivos y actividades.",
    )
    add_heading(doc, "5.7 Experiencia Visual Y Navegación", 2)
    add_body(
        doc,
        "La interfaz incorpora modo claro y oscuro, navegación responsive, estados vacíos con acciones de recuperación y un footer global. El pie de página aporta identidad, accesos internos, redes sociales, ubicación y una recomendación para verificar las condiciones de cada recorrido.",
    )

    add_heading(doc, "6. Modelo De Dominio", 1)
    add_label_table(
        doc,
        [
            ("Atractivo", "Lugar turístico con nombre, departamento, descripción, imagen, URL de YouTube, URL de Google Maps, actividades y circuito asociado cuando corresponde."),
            ("Circuito", "Propuesta de recorrido con nombre, descripción y lista de atractivos relacionados."),
            ("Actividad", "Experiencia concreta con nombre, descripción, duración estimada y atractivo asociado. El costo existe como campo opcional del modelo, pero no se usa en formularios."),
            ("Usuario", "Persona registrada con credenciales, rol y datos necesarios para autenticación e interacción con el sitio."),
        ],
    )

    add_heading(doc, "7. Datos Turísticos Y Cobertura", 1)
    add_body(
        doc,
        "La base de datos fue ampliada para cubrir al menos un atractivo en los 16 departamentos de Catamarca. Además, se agregaron atractivos naturales, arqueológicos, religiosos y museísticos para enriquecer el catálogo.",
    )
    for item in [
        "Naturales: Campo de Piedra Pómez, Volcán Galán, Dunas de Tatón, Cuesta del Portezuelo y Dique de Collagasta.",
        "Arqueológicos y culturales: El Shincal de Quimivil, Parque Arqueológico La Tunita y Santa María del Yokavil.",
        "Religiosos: Catedral Basílica Nuestra Señora del Valle, Gruta de la Virgen del Valle, Monumento a la Virgen del Valle, Monumento a Nuestra Señora de Belén e Iglesia de San Pablo.",
        "Museos: Museo Arqueológico Adán Quiroga, Museo de la Virgen del Valle y Museo Arqueológico Cóndor Huasi.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "8. Arquitectura Técnica", 1)
    add_label_table(
        doc,
        [
            ("Frontend", "Next.js 16.2.6 con App Router, React 19.2.4, componentes cliente, Tailwind CSS 4 y navegación con next/link y next/navigation."),
            ("Backend", "Route Handlers dentro de app/api para atractivos, circuitos, actividades, autenticación, itinerario e informes Excel."),
            ("Base de datos", "MongoDB con Mongoose 9.6.2 y modelos para Atractivo, Circuito, Actividad y Usuario."),
            ("Autenticación", "bcryptjs para hash de contraseñas y jsonwebtoken para tokens JWT."),
            ("Validación", "Helper central en lib/serverValidation.js para validar textos, URLs, IDs, duplicados y referencias."),
            ("SEO", "Metadata global con título, descripción, keywords, Open Graph, robots e idioma español."),
            ("Interfaz", "Tailwind CSS 4, React Icons, tema claro/oscuro, componentes responsive, footer global y estados visuales de carga, error y ausencia de resultados."),
        ],
    )

    add_heading(doc, "9. Reglas De Calidad Aplicadas", 1)
    for item in [
        "Validar formularios en backend, no solo en frontend.",
        "Mantener relaciones consistentes entre atractivos, actividades y circuitos.",
        "Usar imágenes reales cuando sea posible y marcar como generadas las provisorias.",
        "Evitar imágenes de mapas cuando corresponde mostrar un paisaje, monumento o edificio.",
        "Corregir textos con acentos, eñes y nombres propios en UTF-8.",
        "Preservar comportamiento responsive en menús, cards, listas, modales y formularios.",
        "Confirmar borrados y mostrar mensajes de carga, éxito o error.",
        "Ejecutar npm run lint y npm run build para validar estabilidad.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "10. Decisiones Importantes Del Proyecto", 1)
    add_body(
        doc,
        "Se decidió que las actividades dependan del atractivo y no del circuito, porque una experiencia concreta se realiza en un lugar específico. Los circuitos agrupan atractivos y heredan indirectamente las actividades de esos atractivos.",
    )
    add_body(
        doc,
        "La paginación se resolvió desde backend para evitar cargar todos los registros en el cliente y para mantener una navegación estable cuando el catálogo crece.",
    )
    add_body(
        doc,
        "Las imágenes externas se manejan con fallback porque muchas fuentes públicas cambian, bloquean hotlinking o entregan páginas HTML en vez de archivos de imagen. Cuando no se encuentra una foto real estable, se usa una imagen generada marcada explícitamente.",
    )
    add_body(
        doc,
        "La validación se consolidó en backend para que las reglas no dependan solamente del formulario. Así se verifican campos obligatorios, IDs, referencias existentes y URLs incluso cuando la API recibe datos desde otra herramienta.",
    )

    add_heading(doc, "11. Aprendizajes Técnicos", 1)
    for item in [
        "Una relación debe representar el dominio real: la actividad pertenece al atractivo y el circuito funciona como agrupador.",
        "La paginación backend reduce datos transferidos y obliga a conservar correctamente el estado de navegación.",
        "MongoDB acepta caracteres españoles; los problemas de tildes y eñes suelen provenir de una codificación incorrecta.",
        "El HTML inicial debe ser determinista para evitar errores de hidratación al aplicar el tema del navegador.",
        "Los recursos externos requieren fallbacks y revisión de correspondencia, no solamente una URL válida.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "12. Limitaciones Actuales", 1)
    for item in [
        "Algunos videos son búsquedas específicas de YouTube en lugar de enlaces directos confirmados.",
        "Algunas imágenes dependen de servicios externos y pueden fallar si la fuente cambia.",
        "No hay pruebas automatizadas dedicadas.",
        "No existe carga propia de imágenes desde el panel administrativo.",
        "El itinerario se apoya principalmente en estado local del navegador.",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "13. Mejoras Futuras", 1)
    for item in [
        "Incorporar carga de imágenes a Cloudinary, S3 u otro almacenamiento.",
        "Agregar categorías de atractivos y actividades.",
        "Añadir dificultad, temporada recomendada y preparación sugerida.",
        "Sincronizar itinerario por usuario en MongoDB.",
        "Crear pruebas automatizadas para APIs, formularios y flujos críticos.",
        "Mejorar el panel administrador con tablas, búsqueda y edición más cómoda.",
    ]:
        add_bullet(doc, item)

    doc.add_section(WD_SECTION.NEW_PAGE)
    add_heading(doc, "Anexo A. Rutas Y Archivos Relevantes", 1)
    add_label_table(
        doc,
        [
            ("Atractivos", "app/atractivos/page.js, app/atractivos/[id]/page.js, app/api/atractivos/route.js, app/api/atractivos/[id]/route.js"),
            ("Circuitos", "app/circuitos/page.js, app/circuitos/[id]/page.js, app/api/circuitos/route.js, app/api/circuitos/[id]/route.js"),
            ("Actividades", "app/actividades/page.js, app/api/actividades/route.js, app/api/actividades/[id]/route.js"),
            ("Autenticación", "app/login/page.js, app/register/page.js, app/api/auth/*, lib/authMiddleware.js, proxy.js"),
            ("Informes", "app/soporte/page.js, app/api/informes/excel/route.js"),
            ("Componentes", "components/Navbar.js, components/Footer.tsx, components/CardAtractivo.js, components/Toast.js, components/LoadingState.js"),
            ("Modelos", "models/Atractivo.js, models/Circuito.js, models/Actividad.js, models/User.js"),
        ],
    )

    add_heading(doc, "Anexo B. Comandos De Desarrollo", 1)
    add_label_table(
        doc,
        [
            ("Instalar dependencias", "npm install"),
            ("Servidor local", "npm run dev"),
            ("Lint", "npm run lint"),
            ("Build", "npm run build"),
            ("Producción local", "npm run start"),
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
